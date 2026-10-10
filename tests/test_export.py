"""Task 9 / spec §4.2–4.3, §5: authority export, the export:// resource, admin edit."""

from __future__ import annotations

import base64
import hashlib
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def export_dir(tmp_path: Path, monkeypatch):
    d = tmp_path / "exports"
    monkeypatch.setenv("EXPORT_DIR", str(d))
    monkeypatch.setenv("EXPORT_TTL_MINUTES", "60")
    return d


@pytest.mark.anyio
async def test_mixed_tokens_export_only_approved_rows(client, app, planted, export_dir):
    """Review Focus 4: PENDING tokens are skipped and listed, never exported."""
    approved = planted["tokens"]["APPROVED"][:2]
    pending = planted["tokens"]["PENDING_REVIEW"][0]
    r = (await client.call_tool("export_authority_report",
                                {"tokens": approved + [pending, "ghost"], "format": "xlsx"})).structured_content
    assert r["rows"] == 2 and set(r["skipped"]) == {pending, "ghost"} and r["format"] == "xlsx"
    assert r["resource_uri"] == f"export://xlsx/{r['export_id']}"
    path = export_dir / f"{r['export_id']}.xlsx"
    data = path.read_bytes()
    assert hashlib.sha256(data).hexdigest() == r["sha256"]
    ws = load_workbook(io.BytesIO(data)).active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == [label for _, label in app.export.FIELDS]
    assert len(rows) == 3
    names = {app.repo.get_registration(t)["last_name"] for t in approved}
    assert {row[1] for row in rows[1:]} == names
    pending_name = app.repo.get_registration(pending)["last_name"]
    assert pending_name not in {row[1] for row in rows[1:]}
    audit = [a for a in app.repo.audit_recent(20) if a["action"] == "export"]
    assert len(audit) == 1 and r["export_id"] in audit[0]["detail"]
    for n in names:
        assert n not in audit[0]["detail"]


@pytest.mark.anyio
async def test_export_resource_returns_the_same_bytes(client, planted):
    r = (await client.call_tool("export_authority_report",
                                {"tokens": planted["tokens"]["APPROVED"][:3], "format": "xlsx"})).structured_content
    res = await client.read_resource(r["resource_uri"])
    blob = res.contents[0]
    assert blob.mime_type == XLSX
    data = base64.b64decode(blob.blob)
    assert hashlib.sha256(data).hexdigest() == r["sha256"]
    with pytest.raises(Exception):
        await client.read_resource("export://xlsx/doesnotexist")


@pytest.mark.anyio
async def test_window_export_as_pdf(client):
    r = (await client.call_tool("export_authority_report",
                                {"arrival_from": "2026-09-01", "arrival_to": "2026-09-30", "format": "pdf"})).structured_content
    assert r["rows"] == 15 and r["skipped"] == [] and r["resource_uri"].startswith("export://pdf/")
    res = await client.read_resource(r["resource_uri"])
    assert res.contents[0].mime_type == "application/pdf"
    assert base64.b64decode(res.contents[0].blob)[:5] == b"%PDF-"


@pytest.mark.anyio
async def test_export_needs_a_selection_and_rows(client, planted):
    r = await client.call_tool("export_authority_report", {"format": "xlsx"})
    assert r.is_error and "tokens" in r.content[0].text
    r = await client.call_tool("export_authority_report", {"tokens": planted["tokens"]["DRAFT"][:2], "format": "xlsx"})
    assert r.is_error and "no approved" in r.content[0].text.lower()


@pytest.mark.anyio
async def test_export_and_resource_refused_without_pii(client_with_scopes, planted):
    async with client_with_scopes({"export", "pii:read"}) as full:
        r = (await full.call_tool("export_authority_report",
                                  {"tokens": planted["tokens"]["APPROVED"][:1], "format": "xlsx"})).structured_content
    async with client_with_scopes({"export"}) as c:
        denied = await c.call_tool("export_authority_report", {"tokens": planted["tokens"]["APPROVED"][:1], "format": "xlsx"})
        assert denied.is_error and "pii:read" in denied.content[0].text
        with pytest.raises(Exception) as exc:
            await c.read_resource(r["resource_uri"])
        assert "pii:read" in str(exc.value)


@pytest.mark.anyio
async def test_expired_exports_are_swept(client, planted, export_dir, monkeypatch):
    import os
    import time

    r1 = (await client.call_tool("export_authority_report",
                                 {"tokens": planted["tokens"]["APPROVED"][:1], "format": "xlsx"})).structured_content
    old = export_dir / f"{r1['export_id']}.xlsx"
    past = time.time() - 2 * 3600
    os.utime(old, (past, past))
    await client.call_tool("export_authority_report", {"tokens": planted["tokens"]["APPROVED"][:1], "format": "xlsx"})
    assert not old.exists()
    with pytest.raises(Exception):
        await client.read_resource(r1["resource_uri"])


@pytest.mark.anyio
async def test_edit_approved_changes_only_named_fields_and_audits_names_not_values(client, app, planted):
    token = planted["tokens"]["APPROVED"][12]
    before = app.repo.get_registration(token)
    r = (await client.call_tool("edit_approved_registration",
                                {"token": token, "fields": {"car_reg": "7B2 4411", "visa_no": "V000999"}})).structured_content
    assert sorted(r["changed"]) == ["car_reg", "visa_no"] and r["serial"] == r["old_serial"] == before["serial"]
    after = app.repo.get_registration(token)
    assert after["car_reg"] == "7B2 4411" and after["visa_no"] == "V000999"
    for f in ("last_name", "first_name", "dob", "passport_no", "addr_display", "arrival"):
        assert after[f] == before[f], f
    audit = [a for a in app.repo.audit_recent(20) if a["action"] == "edit" and a["actor"] == "mcp:stdio"]
    assert len(audit) == 1 and "car_reg" in audit[0]["detail"] and "7B2" not in audit[0]["detail"]


@pytest.mark.anyio
async def test_edit_moving_the_arrival_reissues_the_serial(client, app, planted):
    token = planted["tokens"]["APPROVED"][13]
    before = app.repo.get_registration(token)
    r = (await client.call_tool("edit_approved_registration",
                                {"token": token, "fields": {"arrival": "2026-09-29", "departure": "2026-10-01"}})).structured_content
    assert r["serial"] != before["serial"] and r["serial"].startswith("2026-09-29-")
    assert set(r["changed"]) >= {"arrival", "departure"}


@pytest.mark.anyio
async def test_edit_rejects_non_approved_unknown_fields_and_bad_values(client, planted):
    pending = planted["tokens"]["PENDING_REVIEW"][5]
    r = await client.call_tool("edit_approved_registration", {"token": pending, "fields": {"car_reg": "X"}})
    assert r.is_error and "APPROVED" in r.content[0].text
    token = planted["tokens"]["APPROVED"][14]
    r = await client.call_tool("edit_approved_registration", {"token": token, "fields": {"serial": "hack"}})
    assert r.is_error and "serial" in r.content[0].text
    r = await client.call_tool("edit_approved_registration", {"token": token, "fields": {"departure": "2020-01-01"}})
    assert r.is_error and "departure" in r.content[0].text
    r = await client.call_tool("edit_approved_registration", {"token": "ghost", "fields": {"car_reg": "X"}})
    assert r.is_error and "list_approved_registrations" in r.content[0].text


@pytest.mark.anyio
async def test_export_and_edit_annotations(client):
    tools = {t.name: t.annotations for t in (await client.list_tools()).tools}
    e = tools["export_authority_report"]
    assert e.read_only_hint is False and e.destructive_hint is False
    d = tools["edit_approved_registration"]
    assert d.read_only_hint is False and d.destructive_hint is False and d.idempotent_hint is True
